"""Parse and serialize plan Excel workbooks."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook

from app.domain.errors import DomainError, InvalidDependencyError, InvalidDurationError
from app.excel import EXPORT_EXTRA, EXPORT_HEADERS_RU, HEADER_ALIASES
from app.schemas import PlanRead


class ExcelParseError(DomainError):
    code = "INVALID_EXCEL"

    def __init__(self, message: str, **details: Any) -> None:
        super().__init__(message, details=details)


def _normalize_header(value: Any) -> str:
    return str(value or "").strip().lower()


def parse_predecessors(raw: Any) -> list[str]:
    if raw is None:
        return []
    text = str(raw).strip()
    if not text:
        return []
    parts = [p.strip() for p in text.replace(";", ",").split(",")]
    return [p for p in parts if p]


def import_xlsx(content: bytes, *, project_start: date | None = None) -> dict:
    """
    Return { name, project_start, rows: [{name, description, assignee, duration, predecessors}] }.
    """
    if not content:
        raise ExcelParseError("Empty file")
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExcelParseError(f"Cannot open workbook: {exc}") from exc

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ExcelParseError("Workbook is empty") from exc

    mapping: dict[int, str] = {}
    for idx, cell in enumerate(header_row):
        key = HEADER_ALIASES.get(_normalize_header(cell))
        if key:
            mapping[idx] = key

    required = {"name", "description", "assignee", "duration", "predecessors"}
    present = set(mapping.values())
    missing = required - present
    if missing:
        raise ExcelParseError(
            "Missing required columns",
            missing=sorted(missing),
            expected=list(EXPORT_HEADERS_RU),
        )

    parsed_rows: list[dict] = []
    for row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        data: dict[str, Any] = {
            "name": "",
            "description": "",
            "assignee": None,
            "duration": None,
            "predecessors": [],
        }
        for idx, field in mapping.items():
            value = row[idx] if idx < len(row) else None
            if field == "predecessors":
                data[field] = parse_predecessors(value)
            elif field == "assignee":
                text = str(value).strip() if value is not None else ""
                data[field] = text or None
            elif field == "duration":
                data[field] = value
            else:
                data[field] = "" if value is None else str(value).strip()

        if not data["name"]:
            raise ExcelParseError(f"Row {row_num}: empty task name", row=row_num)
        try:
            duration = int(data["duration"])
        except (TypeError, ValueError) as exc:
            raise ExcelParseError(
                f"Row {row_num}: duration must be an integer",
                row=row_num,
            ) from exc
        if duration < 1:
            raise InvalidDurationError(duration)
        data["duration"] = duration
        parsed_rows.append(data)

    if not parsed_rows:
        raise InvalidDependencyError("Import contains no tasks")

    return {
        "name": "Imported Plan",
        "project_start": project_start or date.today(),
        "rows": parsed_rows,
    }


def export_xlsx(plan: PlanRead) -> bytes:
    assignee_by_id = {a.id: a.name for a in plan.assignees}
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws.append(EXPORT_HEADERS_RU + EXPORT_EXTRA)

    for task in plan.tasks:
        pred_names = []
        id_to_name = {t.id: t.name for t in plan.tasks}
        for pid in task.predecessor_ids:
            pred_names.append(id_to_name.get(pid, pid))
        ws.append(
            [
                task.name,
                task.description,
                assignee_by_id.get(task.assignee_id or "", "") if task.assignee_id else "",
                task.duration,
                ", ".join(pred_names),
                task.id,
                task.start_date.isoformat(),
                task.end_date.isoformat(),
            ]
        )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
